"""Open a workbook written by xlsx.js with a real spreadsheet reader.

The writer is hand-rolled, so "it produced bytes" proves nothing — Excel's rules
(zip layout, relationship ids, inline strings, cell refs) have to actually hold.
"""
import sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1])
assert wb.sheetnames == ["Details", "Summary"], wb.sheetnames

d = wb["Details"]
assert d.cell(1, 1).value == "Serial", d.cell(1, 1).value
assert d.cell(2, 1).value == "C02ABC123", d.cell(2, 1).value
assert d.cell(2, 3).value == "B", d.cell(2, 3).value
assert d.cell(2, 6) .value == 91, repr(d.cell(2, 6).value)      # a number stays numeric
assert isinstance(d.cell(2, 6).value, int), type(d.cell(2, 6).value)
assert d.cell(3, 2).value == "شاشة مكسورة", d.cell(3, 2).value   # UTF-8 round trip
assert d.cell(4, 3).value is None, repr(d.cell(4, 3).value)      # blank stays blank
assert d.max_row == 4, d.max_row

s = wb["Summary"]
assert s.cell(1, 1).value == "Metric"
assert s.cell(2, 2).value == 3

print("workbook OK —", wb.sheetnames, "| Details rows:", d.max_row)
